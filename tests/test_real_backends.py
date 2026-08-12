from __future__ import annotations

import time
from dataclasses import replace

from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation

from all2md.config import Settings
from all2md.converter import ConversionBackend, convert_file_with_backend
from all2md.server import create_app


def minimal_pdf(text: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        (
            b"<< /Length 49 >>\nstream\nBT /F1 18 Tf 72 720 Td ("
            + text.encode("ascii")
            + b") Tj ET\nendstream"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    document = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, 1):
        offsets.append(len(document))
        document.extend(f"{number} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(document)
    document.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        document.extend(f"{offset:010d} 00000 n \n".encode())
    document.extend(
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF\n"
    )
    return bytes(document)


def test_real_anydoc_liteparse_and_restart_persistence(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("ALL2MD_LITEPARSE_OCR", "false")
    settings = replace(
        Settings.from_env(),
        data_dir=tmp_path,
        conversion_timeout_seconds=30,
        worker_poll_seconds=0.02,
        cleanup_interval_seconds=60,
        log_json=False,
    )
    app = create_app(settings)

    with TestClient(app) as client:
        csv_response = client.post(
            "/convert/json",
            files={"file": ("scores.csv", b"name,score\nAlice,10\n")},
        )
        pdf_response = client.post(
            "/convert/json",
            files={"file": ("hello.pdf", minimal_pdf("Hello LiteParse"))},
        )
        created = client.post(
            "/convert/jobs",
            headers={"Idempotency-Key": "real-restart-test"},
            files={"file": ("durable.csv", b"name,score\nBob,20\n")},
        ).json()
        job_id = created["job_id"]
        deadline = time.time() + 30
        job = created
        while job["status"] != "completed" and time.time() < deadline:
            time.sleep(0.05)
            job = client.get(f"/convert/jobs/{job_id}").json()

    assert csv_response.status_code == 200
    assert csv_response.json()["backend"] == "anydoc"
    assert "Alice" in csv_response.json()["content"]
    assert pdf_response.status_code == 200
    assert pdf_response.json()["backend"] == "liteparse"
    assert "Hello LiteParse" in pdf_response.json()["content"]
    assert job["status"] == "completed"

    restarted_app = create_app(settings)
    with TestClient(restarted_app) as client:
        recovered = client.get(f"/convert/jobs/{job_id}")
        result = client.get(f"/convert/jobs/{job_id}/result")

    assert recovered.json()["status"] == "completed"
    assert "Bob" in result.text


def test_real_markitdown_all_is_the_final_office_fallback(tmp_path, monkeypatch) -> None:
    def fail_specialized_backend(path):
        raise RuntimeError("specialized backend unavailable")

    monkeypatch.setattr("all2md.converter._convert_with_anydoc", fail_specialized_backend)
    monkeypatch.setattr("all2md.converter._convert_with_docling", fail_specialized_backend)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Metric", "Value"])
    worksheet.append(["Revenue", 42])
    spreadsheet = tmp_path / "report.xlsx"
    workbook.save(spreadsheet)

    spreadsheet_result = convert_file_with_backend(
        spreadsheet,
        tmp_path / "report-xlsx.md",
    )

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly report"
    slide.placeholders[1].text = "Revenue increased"
    slides = tmp_path / "report.pptx"
    presentation.save(slides)

    slides_result = convert_file_with_backend(slides, tmp_path / "report-pptx.md")

    assert spreadsheet_result.backend is ConversionBackend.MARKITDOWN
    assert "Revenue" in spreadsheet_result.content
    assert spreadsheet_result.content.strip()
    assert slides_result.backend is ConversionBackend.MARKITDOWN
    assert "Quarterly report" in slides_result.content
